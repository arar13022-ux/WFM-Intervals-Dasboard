"""
WFM Command Center - Streamlit dashboard
Run:  pip install streamlit pandas numpy openpyxl
      streamlit run wfm_dashboard.py

Modules: 1) Roster & shifts  2) Breaks & buffer health  3) 30-min heatmap & OT  4) Attendance & allowances
Assumption: the weekly roster repeats, so overnight shifts on the last day spill into the first day.
"""
import re
from datetime import datetime
from io import BytesIO

import numpy as np
import pandas as pd
import streamlit as st

st.set_page_config(page_title="WFM Command Center", page_icon="📊", layout="wide")
st.markdown(
    """<style>
div[data-testid="stMetric"]{background:#f6f8fb;border:1px solid #e3e8ef;border-radius:10px;padding:10px 14px}
@media (prefers-color-scheme: dark){div[data-testid="stMetric"]{background:#1e2530;border-color:#303a49}}
</style>""",
    unsafe_allow_html=True,
)

# ----------------------------------------------------------------------------- constants
META = ["Employee ID", "Name", "SF Name", "LOB", "TL"]
SAMPLE_DATES = ["27-Sep", "28-Sep", "29-Sep", "30-Sep", "1-Oct", "2-Oct", "3-Oct"]
SHIFTS = ["14:00 - 23:00", "15:00 - 00:00", "17:00 - 02:00", "19:00 - 04:00",
          "20:00 - 05:00", "21:00 - 06:00", "23:00 - 08:00"]
EXC_TYPES = ["None", "Late Login", "Early Leave", "Unplanned Absence"]
BREAK_COLS = ["First Short Break", "Lunch Break", "Last Short Break"]
EDIT_COLS = BREAK_COLS + ["Exception", "Exception Minutes"]
SLOTS = [f"{h:02d}:{m:02d}" for h in range(24) for m in (0, 30)]
G96 = np.arange(96) * 30  # slot start minutes over 2 days (shift day + next day)
UNIT_F = {"Minutes": 30, "Hours": 0.5, "FTE": 1}  # 1 FTE = 30 min in a 30-min interval
UNIT_SFX = {"Minutes": "(min)", "Hours": "(hrs)", "FTE": "(FTE)"}
UNIT_FMT = {"Minutes": "{:.0f}", "Hours": "{:.2f}", "FTE": "{:.2f}"}
# Tunable settings (overridden from the sidebar on every run)
YEAR, OPT_THR, LATE_H, NIGHT_H = 2026, 1.0, 15, 19
EARLIEST, LATEST, GAP, MIN_REST, MAX_CONSEC, VTO_KEEP = 60, 60, 60, 8, 6, 2.0
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHIFT_COLORS = ["#dbeafe", "#e0e7ff", "#ede9fe", "#fce7f3", "#ffedd5", "#fef9c3", "#d1fae5"]


# ----------------------------------------------------------------------------- helpers
def to_min(t):
    try:
        h, m = str(t).strip().split(":")[:2]
        return int(h) * 60 + int(m)
    except Exception:
        return None


def fmt(m):
    return f"{(int(m) // 60) % 24:02d}:{int(m) % 60:02d}"


def parse_shift(s):
    """'17:00 - 02:00' -> (1020, 1560). Returns None for OFF/invalid. End <= start means overnight."""
    if not isinstance(s, str) or "-" not in s:
        return None
    a, b = s.split("-", 1)
    sm, em = to_min(a), to_min(b)
    if sm is None or em is None:
        return None
    return sm, (em + 1440 if em <= sm else em)


def clean_shift(x):
    s = str(x).strip()
    return "OFF" if s.upper() in ("OFF", "", "NAN", "NONE") else re.sub(r"\s*-\s*", " - ", s)


def norm_col(c):
    return f"{c.day}-{c.strftime('%b')}" if hasattr(c, "strftime") else str(c).strip()


def frac(a0, a1):
    """Fraction (0-1) of each 30-min slot in the 96-slot span covered by [a0, a1) minutes."""
    return np.clip(np.minimum(G96 + 30, a1) - np.maximum(G96, a0), 0, 30) / 30


def night_tier(shift):
    p = parse_shift(shift)
    if not p:
        return "None"
    h = p[0] // 60
    return "Overnight" if h >= NIGHT_H else "Late" if h >= LATE_H else "None"


def smap(styler, fn, **kw):
    return (styler.map if hasattr(styler, "map") else styler.applymap)(fn, **kw)


def read_any(f):
    return pd.read_excel(f) if f.name.lower().endswith(("xlsx", "xls")) else pd.read_csv(f)


# ----------------------------------------------------------------------------- core engine
def build_vectors(roster, breaks, dates, lunch_min, short_min):
    """Per shift-day array (agents x [sched, absent, lunch, short] x 96 half-hour slots)."""
    out = {}
    for d in dates:
        bk = breaks.get(d)
        bmap = bk.set_index("Employee ID").to_dict("index") if bk is not None and len(bk) else {}
        V = np.zeros((len(roster), 4, 96))
        for i, (eid, sh) in enumerate(zip(roster["Employee ID"], roster[d])):
            p = parse_shift(sh)
            if not p:
                continue
            s, e = p
            sched = frac(s, e)
            r = bmap.get(eid, {})
            exc = r.get("Exception", "None")
            mins = pd.to_numeric(r.get("Exception Minutes", 0), errors="coerce")
            mins = 0 if pd.isna(mins) else float(mins)
            if exc == "Unplanned Absence":
                ab = sched.copy()
            elif exc == "Late Login":
                ab = frac(s, s + mins)
            elif exc == "Early Leave":
                ab = frac(e - mins, e)
            else:
                ab = np.zeros(96)
            ab = np.minimum(ab, sched)
            present = sched - ab

            def bstart(t):
                m = to_min(t)
                return None if m is None else (m + 1440 if m < s else m)

            lu, sb = np.zeros(96), np.zeros(96)
            m = bstart(r.get("Lunch Break"))
            if m is not None:
                lu = frac(m, m + lunch_min)
            for c in ("First Short Break", "Last Short Break"):
                m = bstart(r.get(c))
                if m is not None:
                    sb += frac(m, m + short_min)
            lu = np.minimum(lu, present)
            sb = np.minimum(sb, np.clip(present - lu, 0, None))
            V[i] = [sched, ab, lu, sb]
        out[d] = V
    return out


def agg(vecs, dates, di, mask_fn):
    """Sum of (sched, absent, lunch, short) per 30-min slot of day `di`, incl. spill-over from previous day."""
    cur = vecs[dates[di]][mask_fn(dates[di])][:, :, :48].sum(axis=0)
    prv = dates[di - 1]  # di=0 wraps to last day (weekly pattern repeats)
    return cur + vecs[prv][mask_fn(prv)][:, :, 48:].sum(axis=0)


def tgt_for(targets, lobs, d=None):
    """Client target per slot for the LOBs. Optional `Date` (e.g. 27-Sep) or `Day` (Mon..Sun) column = per-day targets."""
    t = targets
    if d is not None:
        try:
            if "Date" in t.columns:
                sub = t[t["Date"].map(norm_col) == d]
            elif "Day" in t.columns:
                dow = pd.to_datetime(f"{d}-{YEAR}", format="%d-%b-%Y").strftime("%a")
                sub = t[t["Day"].astype(str).str.strip().str[:3].str.title() == dow]
            else:
                sub = t
            t = sub if len(sub) else t
        except Exception:
            pass
    t = t.set_index("Interval")
    cols = [c for c in lobs if c in t.columns]
    if not cols:
        return np.zeros(48)
    return t[cols].apply(pd.to_numeric, errors="coerce").fillna(0).sum(axis=1).reindex(SLOTS).fillna(0).to_numpy()


def basis_fte(a, basis):
    sched, ab, lu, sb = a
    return sched if basis == "Scheduled FTE" else sched - lu - sb - ab


def status(buf, tgt, sched):
    if tgt == 0 and sched == 0:
        return "⚪ Closed"
    return "🟢 Optimal Buffer" if buf > OPT_THR else "🟡 Fair Buffer" if buf >= 0 else "🔴 Interval Failure"


def interval_table(a, tgt):
    sched, ab, lu, sb = a
    net = sched - lu - sb - ab
    buf = np.round(net - tgt, 2)
    return pd.DataFrame({
        "Interval": SLOTS, "Scheduled FTE": sched.round(2), "On Lunch": lu.round(2),
        "On Short Break": sb.round(2), "Absent/Late": ab.round(2), "Net Floor FTE": net.round(2),
        "Client Target": tgt, "Buffered FTE": buf,
        "Status": [status(b, t, s) for b, t, s in zip(buf, tgt, sched)],
    })


def heat_color(v):
    if pd.isna(v):
        return "background-color:#f1f1f1;color:#aaa"
    if abs(v) < 1e-9:
        return "background-color:#ffeb9c;color:#7a5c00"
    return "background-color:#c6efce;color:#006100" if v > 0 else "background-color:#ffc7ce;color:#9c0006"


def status_color(v):
    return {"🟢": "background-color:#c6efce;color:#006100", "🟡": "background-color:#ffeb9c;color:#7a5c00",
            "🔴": "background-color:#ffc7ce;color:#9c0006"}.get(str(v)[:1], "color:#999")


# ----------------------------------------------------------------------------- sample data
def default_breaks(roster, d):
    rows = []
    for i, r in roster.reset_index(drop=True).iterrows():
        p = parse_shift(r[d])
        if not p:
            continue
        s, k = p[0], i % 4
        rows.append({"Employee ID": r["Employee ID"], "Name": r["Name"], "LOB": r["LOB"], "TL": r["TL"],
                     "Shift": r[d], "First Short Break": fmt(s + 120 + 15 * k), "Lunch Break": fmt(s + 240 + 30 * k),
                     "Last Short Break": fmt(s + 420 + 15 * (k % 2)), "Exception": "None", "Exception Minutes": 0})
    return pd.DataFrame(rows)


def init_sample():
    rng = np.random.default_rng(7)
    first = ["Omar", "Mona", "Youssef", "Salma", "Hady", "Nour", "Karim", "Dina", "Mostafa", "Heba", "Tarek", "Rana",
             "Amr", "Yara", "Khaled", "Laila", "Ziad", "Farida", "Mahmoud", "Aya", "Sherif", "Mariam", "Hossam",
             "Nada", "Ibrahim", "Jana", "Walid", "Reem"]
    last = ["Fathy", "Saad", "Nabil", "Gamal", "Mansour", "Ashraf", "Helmy", "Zaki", "Farouk", "Lotfy"]
    rows = []
    for i in range(28):
        eng = i < 16
        row = {"Employee ID": f"E{1001 + i}", "Name": f"{first[i]} {last[i % 10]}",
               "SF Name": f"{first[i].lower()}.{last[i % 10].lower()}", "LOB": "NMG-ENG" if eng else "NMG-SP",
               "TL": ["Ahmed Hassan", "Mona Samir"][i % 2] if eng else "Karim Adel"}
        off = int(rng.integers(0, 7))
        for j, d in enumerate(SAMPLE_DATES):
            if j in (off, (off + 1) % 7):
                row[d] = "OFF"
            else:
                row[d] = SHIFTS[i % 7] if rng.random() > 0.15 else SHIFTS[int(rng.integers(0, 7))]
        rows.append(row)
    roster = pd.DataFrame(rows)
    breaks = {}
    for d in SAMPLE_DATES:
        b = default_breaks(roster, d)
        for idx in b.index:
            x = rng.random()
            if x < 0.05:
                b.loc[idx, "Exception"] = "Unplanned Absence"
            elif x < 0.10:
                b.loc[idx, ["Exception", "Exception Minutes"]] = ["Late Login", int(rng.choice([15, 30, 60]))]
            elif x < 0.13:
                b.loc[idx, ["Exception", "Exception Minutes"]] = ["Early Leave", 60]
        breaks[d] = b
    vecs = build_vectors(roster, breaks, SAMPLE_DATES, 30, 15)
    tg = {"Interval": SLOTS}
    for lob in ("NMG-ENG", "NMG-SP"):
        acc = sum(agg(vecs, SAMPLE_DATES, di, lambda d, lob=lob: (roster["LOB"] == lob).to_numpy())[0]
                  for di in range(7)) / 7
        tg[lob] = np.floor(acc * 0.8 + 0.3).astype(int)
    st.session_state.update(roster=roster, breaks=breaks, targets=pd.DataFrame(tg))


def prep_roster(r):
    r = r.copy()
    r.columns = [norm_col(c) for c in r.columns]
    for c in META:
        if c not in r:
            r[c] = ""
    dates = [c for c in r.columns if c not in META]
    for d in dates:
        r[d] = r[d].map(clean_shift)
    return r[META + dates]


def apply_uploads(f_r, f_t, f_b, f_s=None):
    if f_s:  # full session restore (roster + targets + breaks + change log)
        sh = pd.read_excel(f_s, sheet_name=None)
        r = prep_roster(sh["Roster"])
        t = sh["Targets"]
        t["Interval"] = t["Interval"].astype(str).str[:5]
        bk = sh["Breaks"]
        bk["Date"] = bk["Date"].map(norm_col)
        dates = [c for c in r.columns if c not in META]
        grp = {d: g.drop(columns="Date").reset_index(drop=True) for d, g in bk.groupby("Date")}
        st.session_state.update(roster=r, targets=t, breaks={d: grp.get(d, default_breaks(r, d)) for d in dates})
        lg = sh.get("Change log")
        st.session_state.log = lg.to_dict("records") if lg is not None and len(lg) else []
    if f_r:
        r = prep_roster(read_any(f_r))
        st.session_state.roster = r
        st.session_state.breaks = {d: default_breaks(r, d) for d in r.columns if d not in META}
    if f_t:
        t = read_any(f_t)
        t["Interval"] = t["Interval"].astype(str).str[:5]
        st.session_state.targets = t
    if f_b:
        b = read_any(f_b)
        b["Date"] = b["Date"].map(norm_col)
        for d, g in b.groupby("Date"):
            if d in st.session_state.breaks:
                base = st.session_state.breaks[d].set_index("Employee ID")
                base.update(g.set_index("Employee ID")[[c for c in EDIT_COLS if c in g]])
                st.session_state.breaks[d] = base.reset_index()


if "roster" not in st.session_state:
    init_sample()

# ----------------------------------------------------------------------------- sidebar
with st.sidebar:
    st.title("📊 WFM Command Center")
    with st.expander("📂 Data inputs (optional)"):
        f_r = st.file_uploader("Master roster (CSV/XLSX)", type=["csv", "xlsx"], key="u_r")
        f_t = st.file_uploader("Client FTE targets: Interval + one column per LOB", type=["csv", "xlsx"], key="u_t")
        f_b = st.file_uploader("Break/exception log: Date, Employee ID, breaks, Exception, Exception Minutes",
                               type=["csv", "xlsx"], key="u_b")
        f_s = st.file_uploader("Restore saved session (.xlsx from 'Save session')", type=["xlsx"], key="u_s")
        c1, c2 = st.columns(2)
        if c1.button("Apply"):
            apply_uploads(f_r, f_t, f_b, f_s)
            st.rerun()
        if c2.button("Reset sample"):
            init_sample()
            st.rerun()

    roster, targets, breaks = st.session_state.roster, st.session_state.targets, st.session_state.breaks
    DATES = [c for c in roster.columns if c not in META]
    all_shifts = sorted({s for d in DATES for s in roster[d].unique()},
                        key=lambda s: (SHIFTS.index(s) if s in SHIFTS else 99, s))

    st.header("🎛️ Global filters")
    sel_date = st.selectbox("Date", DATES)
    sel_lobs = st.multiselect("LOB", sorted(roster["LOB"].unique()), default=sorted(roster["LOB"].unique()))
    sel_tls = st.multiselect("Team Leader", sorted(roster["TL"].unique()), default=sorted(roster["TL"].unique()))
    sel_shifts = st.multiselect("Shift type", all_shifts, default=all_shifts)

    st.header("⚙️ Settings")
    lunch_min = st.select_slider("Lunch duration (min)", [30, 60], 30)
    short_min = st.select_slider("Short break duration (min)", [15, 30], 15)
    basis = st.radio("Heatmap/OT basis", ["Scheduled FTE", "Net Floor FTE"], help="Net = after breaks and exceptions")
    unit = st.radio("Display unit", ["Minutes", "Hours", "FTE"], help="1 FTE = 30 minutes of staffed time per 30-min interval")
    OPT_THR = st.number_input("Optimal buffer above (FTE)", 0.0, 10.0, 1.0, 0.5,
                              help="🟢 buffer above this · 🟡 from 0 up to it · 🔴 below 0")
    late_rate = st.number_input("Late-shift allowance (EGP)", 0, 1000, 75, 5)
    night_rate = st.number_input("Overnight allowance (EGP)", 0, 1000, 150, 5)
    LATE_H = st.number_input("Late allowance starts at shift-start hour", 0, 23, 15)
    NIGHT_H = st.number_input("Overnight allowance starts at shift-start hour", 0, 23, 19)
    YEAR = st.number_input("Roster year (for per-weekday targets)", 2020, 2040, 2026)
    with st.expander("Break & roster rules"):
        EARLIEST = st.number_input("Earliest break after shift start (min)", 0, 300, 60, 15)
        LATEST = st.number_input("Last break must end before shift end by (min)", 0, 300, 60, 15)
        GAP = st.number_input("Minimum gap between breaks (min)", 0, 300, 60, 15)
        MIN_REST = st.number_input("Minimum rest between shifts (hours)", 0, 24, 8)
        MAX_CONSEC = st.number_input("Max consecutive working days", 1, 14, 6)
        VTO_KEEP = st.number_input("VTO: keep this much buffer (FTE)", 0.0, 10.0, 2.0, 0.5)


def mask_fn(d, lobs=None):
    lobs = sel_lobs if lobs is None else lobs
    return (roster["LOB"].isin(lobs) & roster["TL"].isin(sel_tls) & roster[d].isin(sel_shifts)).to_numpy()


# ----------------------------------------------------------------------------- v2 helpers
def agent_day(d):
    """One row per working agent on shift-date d, with attendance, lost hours and allowance."""
    ad = roster.loc[mask_fn(d), META + [d]].rename(columns={d: "Shift"})
    ad = ad[ad["Shift"].map(lambda x: parse_shift(x) is not None)]
    if ad.empty:
        return ad
    ad = ad.merge(breaks[d][["Employee ID", "Exception", "Exception Minutes"]], on="Employee ID", how="left")
    ad["Exception"] = ad["Exception"].fillna("None")
    ad["Exception Minutes"] = pd.to_numeric(ad["Exception Minutes"], errors="coerce").fillna(0)
    ad["Hours"] = ad["Shift"].map(lambda x: (parse_shift(x)[1] - parse_shift(x)[0]) / 60)
    ad["Absent"] = (ad["Exception"] == "Unplanned Absence").astype(int)
    ad["Lost"] = np.where(ad["Absent"] == 1, ad["Hours"],
                          np.where(ad["Exception"].isin(["Late Login", "Early Leave"]),
                                   np.minimum(ad["Exception Minutes"] / 60, ad["Hours"]), 0))
    ad["Attended"] = 1 - ad["Absent"]
    ad["Tier"] = ad["Shift"].map(night_tier)
    ad["Late Shifts"] = ((ad["Tier"] == "Late") & (ad["Attended"] == 1)).astype(int)
    ad["Overnight Shifts"] = ((ad["Tier"] == "Overnight") & (ad["Attended"] == 1)).astype(int)
    ad["Night Allowance (EGP)"] = ad["Late Shifts"] * late_rate + ad["Overnight Shifts"] * night_rate
    return ad


def summarize(ad, by):
    g = ad.groupby(by).agg(**{
        "Total Scheduled": ("Employee ID", "count"), "Attended": ("Attended", "sum"),
        "Unplanned Absences": ("Absent", "sum"), "Lost Hours": ("Lost", "sum"), "Sched Hours": ("Hours", "sum"),
        "Late Shifts": ("Late Shifts", "sum"), "Overnight Shifts": ("Overnight Shifts", "sum"),
        "Night Allowance (EGP)": ("Night Allowance (EGP)", "sum")}).reset_index()
    tot = g.drop(columns=by).sum()
    tot[by] = "TOTAL"
    g = pd.concat([g, tot.to_frame().T], ignore_index=True)
    for c in g.columns[1:]:
        g[c] = pd.to_numeric(g[c])
    g["Attendance %"] = (g["Attended"] / g["Total Scheduled"] * 100).round(1)
    g["Shrinkage %"] = (g["Lost Hours"] / g["Sched Hours"] * 100).round(1)
    return g


def validate_breaks(df):
    """Break-rule violations for rows of the editor view (needs 'Shift Date', Shift, break columns)."""
    out = []
    for _, r in df.iterrows():
        p = parse_shift(r["Shift"])
        if not p or r["Exception"] == "Unplanned Absence":
            continue
        s0, e0 = p
        items = []
        for lbl, c, dur in (("First short", "First Short Break", short_min), ("Lunch", "Lunch Break", lunch_min),
                            ("Last short", "Last Short Break", short_min)):
            m = to_min(r[c])
            if m is None:
                out.append((r, f"{lbl} break missing"))
                continue
            m = m + 1440 if m < s0 else m
            items.append((lbl, m, m + dur))
            if m < s0 + EARLIEST:
                out.append((r, f"{lbl} starts {m - s0} min after shift start (min {EARLIEST})"))
            if m + dur > e0 - LATEST:
                out.append((r, f"{lbl} ends {e0 - m - dur} min before shift end (min {LATEST})"))
        for (l1, a1, b1), (l2, a2, b2) in zip(items, items[1:]):
            if a2 - b1 < GAP:
                out.append((r, f"{l1} to {l2}: gap {a2 - b1} min (min {GAP})"))
    return pd.DataFrame([{"Shift Date": r["Shift Date"], "Employee ID": r["Employee ID"], "Name": r["Name"], "Issue": m}
                         for r, m in out])


def auto_balance(di_):
    """Re-time lunch/short breaks of agents starting on DATES[di_] to maximise the buffer within the break rules."""
    n = len(DATES)
    if n < 2:
        return None
    cur, prv, nxt = DATES[di_], DATES[di_ - 1], DATES[(di_ + 1) % n]

    def net(d, lo, hi):
        V = vecs[d][mask_fn(d)][:, :, lo:hi]
        return (V[:, 0] - V[:, 1] - V[:, 2] - V[:, 3]).sum(axis=0)

    buf = np.zeros(96)  # timeline: cur date 00:00 -> next date 24:00
    if prv != cur:
        buf[:48] += net(prv, 48, 96)
    if nxt != cur:
        buf[48:] += net(nxt, 0, 48)
    pos = {e: i for i, e in enumerate(roster["Employee ID"])}
    ids = set(roster.loc[mask_fn(cur), "Employee ID"])
    bk = breaks[cur].copy()
    agents = []
    for _, r in bk.iterrows():
        p = parse_shift(r["Shift"])
        if r["Employee ID"] not in ids or r["Exception"] == "Unplanned Absence" or not p:
            continue
        v = vecs[cur][pos[r["Employee ID"]]]
        pres = v[0] - v[1]
        buf += pres
        agents.append((p[0], r["Employee ID"], p[1], pres))
    buf -= np.concatenate([tgt_for(targets, sel_lobs, cur), tgt_for(targets, sel_lobs, nxt)])
    agents.sort(key=lambda x: (x[0], x[1]))

    def place(lo, hi, dur, pres):
        k, w = int(np.ceil(dur / 30)), min(dur, 30) / 30
        best = None
        for t in range(int(np.ceil(lo / 30)) * 30, int(hi) + 1, 30):
            a = t // 30
            if a + k > 96 or pres[a:a + k].min() < 0.99:
                continue
            sc = (buf[a:a + k].min(), buf[a:a + k].sum())
            if best is None or sc > best[0]:
                best = (sc, t)
        if best is None:
            return None
        buf[best[1] // 30:best[1] // 30 + k] -= w
        return best[1]

    L, S, lunch_t, res = lunch_min, short_min, {}, {}
    for s0, eid, e0, pres in agents:  # 1) lunches
        span = e0 - s0
        lo = max(s0 + span * 0.33, s0 + EARLIEST + S + GAP)
        hi = min(s0 + span * 0.67, e0 - LATEST - S - GAP - L)
        if hi < lo:
            lo, hi = s0 + span * 0.33, s0 + span * 0.67
        t = place(lo, hi, L, pres)
        if t is not None:
            lunch_t[eid] = t
            res.setdefault(eid, {})["Lunch Break"] = fmt(t)
    for s0, eid, e0, pres in agents:  # 2) first short break
        hi = lunch_t[eid] - S - GAP if eid in lunch_t else s0 + (e0 - s0) * 0.33
        t = place(s0 + EARLIEST, hi, S, pres)
        if t is not None:
            res.setdefault(eid, {})["First Short Break"] = fmt(t)
    for s0, eid, e0, pres in agents:  # 3) last short break
        lo = lunch_t[eid] + L + GAP if eid in lunch_t else s0 + (e0 - s0) * 0.67
        t = place(lo, e0 - LATEST - S, S, pres)
        if t is not None:
            res.setdefault(eid, {})["Last Short Break"] = fmt(t)
    idx = bk.set_index("Employee ID")
    for eid, dct in res.items():
        for c, v in dct.items():
            idx.at[eid, c] = v
    return idx.reset_index()


def on_break(k):
    """Agents on lunch / short break in calendar slot k of the selected date."""
    rows = []
    for d, off in ((DATES[di - 1], 48), (sel_date, 0)):
        if off and DATES[di - 1] == sel_date:
            continue
        bk = breaks[d].set_index("Employee ID")
        for i in np.where(mask_fn(d))[0]:
            lu, sb = vecs[d][i, 2, off + k], vecs[d][i, 3, off + k]
            eid = roster["Employee ID"].iat[i]
            if (lu > 0 or sb > 0) and eid in bk.index:
                b = bk.loc[eid]
                rows.append({"Shift Date": d, "Employee ID": eid, "Name": roster["Name"].iat[i],
                             "LOB": roster["LOB"].iat[i], "TL": roster["TL"].iat[i], "Shift": roster[d].iat[i],
                             "Break": " + ".join(n for n, v in (("Lunch", lu), ("Short break", sb)) if v > 0),
                             "First Short": b["First Short Break"], "Lunch": b["Lunch Break"],
                             "Last Short": b["Last Short Break"]})
    return pd.DataFrame(rows)


def windows(flags, min_len=1):
    out, a = [], None
    for k, v in enumerate(list(flags) + [False]):
        if v and a is None:
            a = k
        elif not v and a is not None:
            if k - a >= min_len:
                out.append((a, k))
            a = None
    return out


def shift_edges():
    """(name, shift, start_min, end_min) for shifts touching the selected calendar date (minutes from its midnight)."""
    out = []
    for d, off in ((DATES[di - 1], -1440), (sel_date, 0)):
        if off and DATES[di - 1] == sel_date:
            continue
        for i in np.where(mask_fn(d))[0]:
            p = parse_shift(roster[d].iat[i])
            if p:
                out.append((roster["Name"].iat[i], roster[d].iat[i], p[0] + off, p[1] + off))
    return out


def _names(items, n=6):
    return "-" if not items else ", ".join(items[:n]) + (f" +{len(items) - n}" if len(items) > n else "")


def ot_suggestions(deficit):
    edges = shift_edges()
    offs = roster.loc[(roster[sel_date] == "OFF") & roster["LOB"].isin(sel_lobs) & roster["TL"].isin(sel_tls),
                      "Name"].tolist()
    rows = []
    for a, b in windows(deficit > 0):
        w0, w1 = a * 30, b * 30
        ext = [f"{n} (ends {sh.split(' - ')[1]})" for n, sh, s0, e0 in edges if w0 - 60 <= e0 <= w0]
        early = [f"{n} (starts {sh.split(' - ')[0]})" for n, sh, s0, e0 in edges if w1 <= s0 <= w1 + 60]
        rows.append({"Window": f"{fmt(w0)}-{fmt(w1)}", "Peak deficit (FTE)": round(float(deficit[a:b].max()), 2),
                     "OT hours": round(float(deficit[a:b].sum() * 0.5), 2), "Extend shift": _names(ext),
                     "Start early": _names(early), "OFF-day volunteers": _names(offs)})
    return pd.DataFrame(rows)


def vto_suggestions(buf):
    edges = shift_edges()
    sur = np.clip(buf - VTO_KEEP, 0, None)
    rows = []
    for a, b in windows(sur > 0, 4):
        w0, w1 = a * 30, b * 30
        leave = [f"{n} (ends {sh.split(' - ')[1]})" for n, sh, s0, e0 in edges if w0 < e0 <= w1]
        late = [f"{n} (starts {sh.split(' - ')[0]})" for n, sh, s0, e0 in edges if w0 <= s0 < w1]
        rows.append({"Window": f"{fmt(w0)}-{fmt(w1)}", "Min surplus (FTE)": round(float(sur[a:b].min()), 2),
                     "Max VTO hours": round(float(sur[a:b].sum() * 0.5), 2), "Leave early": _names(leave),
                     "Start late": _names(late)})
    return pd.DataFrame(rows)


def roster_checks(r):
    out = []
    for _, row in r.iterrows():
        run = 0
        for j, d in enumerate(DATES):
            p = parse_shift(row[d])
            run = run + 1 if p else 0
            if run == MAX_CONSEC + 1:
                out.append({"Employee ID": row["Employee ID"], "Name": row["Name"], "Date": d,
                            "Issue": f"More than {MAX_CONSEC} consecutive working days"})
            if p and j + 1 < len(DATES) and parse_shift(row[DATES[j + 1]]):
                rest = (1440 + parse_shift(row[DATES[j + 1]])[0] - p[1]) / 60
                if rest < MIN_REST:
                    out.append({"Employee ID": row["Employee ID"], "Name": row["Name"], "Date": DATES[j + 1],
                                "Issue": f"Only {rest:.1f}h rest after the {d} shift (min {MIN_REST}h)"})
    return pd.DataFrame(out)


def resync_breaks(new_roster, old):
    """Rebuild break tables after roster edits; keep existing break edits for unchanged shifts."""
    out = {}
    for d in DATES:
        fresh = default_breaks(new_roster, d)
        o = old.get(d)
        if o is not None and len(o):
            o = o.set_index("Employee ID")
            for i, r in fresh.iterrows():
                if r["Employee ID"] in o.index and o.at[r["Employee ID"], "Shift"] == r["Shift"]:
                    for c in EDIT_COLS:
                        fresh.at[i, c] = o.at[r["Employee ID"], c]
        out[d] = fresh
    return out


def apply_roster(new, log):
    st.session_state.breaks = resync_breaks(new, st.session_state.breaks)
    st.session_state.roster = new
    st.session_state.log = st.session_state.get("log", []) + log
    st.rerun()


def session_bytes():
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        roster.to_excel(xw, sheet_name="Roster", index=False)
        targets.to_excel(xw, sheet_name="Targets", index=False)
        pd.concat([b.assign(Date=d)[["Date"] + list(b.columns)] for d, b in breaks.items()]).to_excel(
            xw, sheet_name="Breaks", index=False)
        pd.DataFrame(st.session_state.get("log", []), columns=["Time", "Employee ID", "Name", "Date", "From", "To",
                                                                "Source"]).to_excel(xw, sheet_name="Change log", index=False)
    return buf.getvalue()


def make_report():
    from openpyxl.styles import Font, PatternFill
    fills = {"g": PatternFill("solid", fgColor="C6EFCE"), "y": PatternFill("solid", fgColor="FFEB9C"),
             "r": PatternFill("solid", fgColor="FFC7CE"), "n": PatternFill("solid", fgColor="EEEEEE")}
    sheets = {"Roster": view, f"Buffer {sel_date}": itd, "Weekly heatmap": (week * UF).rename_axis("Interval").reset_index(),
              f"OT {sel_date}": ot_df, "Attendance": show, "Exceptions": exc}
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            for c in ws[1]:
                c.font = Font(bold=True)
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 16
            heat_cols = []
            if name == "Weekly heatmap":
                heat_cols = list(range(2, ws.max_column + 1))
            elif name.startswith("Buffer"):
                heat_cols = [i + 1 for i, c in enumerate(df.columns) if c == "Buffered (min)"]
            for ci in heat_cols:
                for row in ws.iter_rows(min_row=2, min_col=ci, max_col=ci):
                    c = row[0]
                    v = c.value
                    c.fill = fills["n"] if v is None else fills["g"] if v > 0 else fills["y"] if v == 0 else fills["r"]
    return buf.getvalue()


vecs = build_vectors(roster, breaks, DATES, lunch_min, short_min)
di = DATES.index(sel_date)
UF = UNIT_F[unit]
tgt = tgt_for(targets, sel_lobs, sel_date)
a_day = agg(vecs, DATES, di, mask_fn)
it = interval_table(a_day, tgt)

st.title("Workforce Management Dashboard")
st.caption(f"Date: **{sel_date}** | LOB: {', '.join(sel_lobs) or '-'} | {len(sel_tls)} TL(s) | {len(sel_shifts)} shift type(s)")
if not sel_lobs:
    st.warning("Select at least one LOB.")
    st.stop()

tab1, tab2, tab3, tab4, tab5 = st.tabs(["📋 Roster & Shifts", "⏱️ Breaks & Buffer Health",
                                        "🔥 Interval Heatmap & OT", "📈 Attendance & Allowances", "👥 Roster Editor"])

# ----------------------------------------------------------------------------- Module 1
with tab1:
    row_mask = roster["LOB"].isin(sel_lobs) & roster["TL"].isin(sel_tls) & roster[sel_date].isin(sel_shifts)
    view = roster[row_mask]
    daily = {"Working": [], "OFF": []}
    for d in DATES:
        col = roster.loc[roster["LOB"].isin(sel_lobs) & roster["TL"].isin(sel_tls), d]
        col = col[col.isin(sel_shifts)]
        daily["Working"].append(int((col != "OFF").sum()))
        daily["OFF"].append(int((col == "OFF").sum()))
    daily = pd.DataFrame(daily, index=DATES)

    k = st.columns(5)
    k[0].metric("Agents shown", len(view))
    k[1].metric("Scheduled shifts (week)", int(daily["Working"].sum()))
    k[2].metric("OFF days (week)", int(daily["OFF"].sum()))
    k[3].metric(f"Working {sel_date}", int(daily.loc[sel_date, "Working"]))
    k[4].metric(f"OFF {sel_date}", int(daily.loc[sel_date, "OFF"]))

    colmap = {s: SHIFT_COLORS[i % 7] for i, s in enumerate(SHIFTS)}
    sty = smap(view.style, lambda v: "background-color:#e5e7eb;color:#9ca3af" if v == "OFF"
               else f"background-color:{colmap.get(v, '#fff')};color:#111", subset=DATES)
    sty = sty.set_properties(subset=[sel_date], **{"border": "2px solid #3b82f6", "font-weight": "600"})
    st.dataframe(sty, hide_index=True, height=min(35 * len(view) + 40, 600))

    c1, c2 = st.columns(2)
    c1.subheader("Working vs OFF per day")
    c1.bar_chart(daily)
    c2.subheader(f"Shift mix on {sel_date}")
    wk = view[view[sel_date] != "OFF"]
    if len(wk):
        mix = pd.crosstab(wk[sel_date], wk["LOB"]).reindex([s for s in SHIFTS if s in set(wk[sel_date])] +
                                                            sorted(set(wk[sel_date]) - set(SHIFTS)))
        c2.bar_chart(mix)
    else:
        c2.info("No working agents for this selection.")

# ----------------------------------------------------------------------------- Module 2
with tab2:
    with st.expander("🪄 Auto-balance breaks", expanded=False):
        st.caption("Re-times lunches and short breaks of agents starting on this date to maximise the buffer, "
                   "within the break rules in the sidebar. Review the result, then accept or discard.")
        if st.button("Generate proposal"):
            st.session_state.prop = (sel_date, auto_balance(di))
        pr = st.session_state.get("prop")
        if pr and pr[0] == sel_date and pr[1] is not None:
            v2 = build_vectors(roster, {**breaks, sel_date: pr[1]}, DATES, lunch_min, short_min)
            it2 = interval_table(agg(v2, DATES, di, mask_fn), tgt)
            fail = lambda x: int((x["Status"] == "🔴 Interval Failure").sum())
            moved = int((pr[1][BREAK_COLS].astype(str).values != breaks[sel_date][BREAK_COLS].astype(str).values).sum())
            m = st.columns(4)
            m[0].metric("Failed intervals", fail(it2), delta=fail(it2) - fail(it), delta_color="inverse")
            m[1].metric("Worst buffer (min)", f"{it2['Buffered FTE'].min() * 30:.0f}",
                        delta=f"{(it2['Buffered FTE'].min() - it['Buffered FTE'].min()) * 30:+.0f}")
            m[2].metric("Peak lunch concurrency", f"{it2['On Lunch'].max():.0f}",
                        delta=f"{it2['On Lunch'].max() - it['On Lunch'].max():+.0f}", delta_color="inverse")
            m[3].metric("Breaks re-timed", moved)
            c1, c2 = st.columns(2)
            if c1.button("✅ Accept proposal"):
                st.session_state.breaks[sel_date] = pr[1]
                st.session_state.prop = None
                st.rerun()
            if c2.button("✖ Discard"):
                st.session_state.prop = None
                st.rerun()

    st.subheader(f"Break schedule & exceptions - {sel_date}")
    base = breaks[sel_date]
    prev = DATES[di - 1]

    def _rows(d, overnight_only=False):
        b = breaks[d]
        b = b[b["Employee ID"].isin(roster.loc[mask_fn(d), "Employee ID"])]
        if overnight_only:  # shifts that run past midnight and so affect the early hours of the next date
            b = b[b["Shift"].map(lambda x: (parse_shift(x) or (0, 0))[1] > 1440)]
        return b.assign(**{"Shift Date": d})

    parts = ([_rows(prev, True)] if prev != sel_date else []) + [_rows(sel_date)]
    bview = pd.concat(parts).reset_index(drop=True)
    bview = bview[["Shift Date", "Employee ID", "Name", "LOB", "TL", "Shift"] + EDIT_COLS]
    st.caption(f"Rows tagged **{prev}** are overnight agents whose breaks after 00:00 fall on **{sel_date}**. "
               f"Breaks after 00:00 for agents tagged **{sel_date}** show up on the next date.")
    sig = hash((sel_date, tuple(sel_lobs), tuple(sel_tls), tuple(sel_shifts)))
    time_cfg = lambda lbl: st.column_config.TextColumn(lbl, validate=r"^([01]\d|2[0-3]):[0-5]\d$", help="HH:MM")
    edited = st.data_editor(
        bview, key=f"ed_{sel_date}_{sig}", hide_index=True, height=min(35 * len(bview) + 40, 420),
        disabled=["Shift Date", "Employee ID", "Name", "LOB", "TL", "Shift"],
        column_config={
            "First Short Break": time_cfg("First Short Break"), "Lunch Break": time_cfg("Lunch Break"),
            "Last Short Break": time_cfg("Last Short Break"),
            "Exception": st.column_config.SelectboxColumn("Exception", options=EXC_TYPES, required=True),
            "Exception Minutes": st.column_config.NumberColumn(
                "Exception Minutes", min_value=0, max_value=540, step=5, help="Minutes late / minutes left early"),
        },
    )

    def _norm(df):
        x = df[EDIT_COLS].copy()
        x["Exception Minutes"] = pd.to_numeric(x["Exception Minutes"], errors="coerce").fillna(0).astype(int)
        return x.fillna("").astype(str).reset_index(drop=True)

    if not _norm(edited).equals(_norm(bview)):  # write edits back to each row's own shift date, then recompute
        for d, g in edited.groupby("Shift Date"):
            new = breaks[d].set_index("Employee ID")
            ed = g.set_index("Employee ID")
            for c in EDIT_COLS:
                new.loc[ed.index, c] = ed[c]
            new["Exception Minutes"] = pd.to_numeric(new["Exception Minutes"], errors="coerce").fillna(0).astype(int)
            st.session_state.breaks[d] = new.reset_index()
        st.rerun()

    viol = validate_breaks(edited)
    if len(viol):
        with st.expander(f"⚠️ {len(viol)} break-rule violation(s)", expanded=False):
            st.dataframe(viol, hide_index=True)
    else:
        st.success("All breaks respect the break rules.")

    st.subheader("Live buffer health")
    cnt = it["Status"].value_counts()
    peak = it.loc[it["On Lunch"].idxmax()]
    k = st.columns(4)
    k[0].metric("🟢 Optimal intervals", int(cnt.get("🟢 Optimal Buffer", 0)))
    k[1].metric("🟡 Fair intervals", int(cnt.get("🟡 Fair Buffer", 0)))
    k[2].metric("🔴 Failed intervals", int(cnt.get("🔴 Interval Failure", 0)))
    k[3].metric("Peak lunch concurrency", f"{peak['On Lunch']:.0f}", help=f"At {peak['Interval']}")

    risk = it[(it["On Lunch"] >= 2) & (it["Buffered FTE"] <= OPT_THR)]
    if len(risk):
        st.warning(f"⚠️ 2+ agents on lunch while buffer ≤ {OPT_THR:g} FTE at: {', '.join(risk['Interval'])}. Consider staggering lunches.")
    else:
        st.success("No lunch-concurrency risk in low-coverage windows.")

    fte_cols = ["Scheduled FTE", "On Lunch", "On Short Break", "Absent/Late", "Net Floor FTE", "Client Target"]
    itd = it.copy()
    itd["Buffered (min)"] = (it["Buffered FTE"] * 30).round(0)  # 1 FTE = 30 min per 30-min interval
    itd = itd.drop(columns="Buffered FTE")[fte_cols + ["Buffered (min)", "Status"]].copy()
    itd.insert(0, "Interval", it["Interval"])
    sty = smap(itd.style, status_color, subset=["Status"])
    sty = smap(sty, heat_color, subset=["Buffered (min)"])
    sty = sty.format({**{c: "{:.2f}" for c in fte_cols}, "Buffered (min)": "{:+.0f}"})
    st.dataframe(sty, hide_index=True, height=520)
    st.line_chart(itd.set_index("Interval")[["Net Floor FTE", "Client Target"]])

    d1, d2 = st.columns(2)
    d1.download_button("⬇️ Interval buffer report (CSV)", itd.to_csv(index=False).encode("utf-8-sig"),
                       f"buffer_report_{sel_date}.csv", "text/csv")
    d2.download_button("⬇️ Break schedule & exceptions (CSV)", base.to_csv(index=False).encode("utf-8-sig"),
                       f"break_schedule_{sel_date}.csv", "text/csv")

    st.subheader("🔎 Who's on break at an interval")
    pick = st.selectbox("Interval", SLOTS, index=int(it["Buffered FTE"].idxmin()), key=f"pick_{sel_date}",
                        help="Defaults to the weakest interval of the day")
    kk = SLOTS.index(pick)
    rw = it.iloc[kk]
    m = st.columns(4)
    m[0].metric("Net floor FTE", f"{rw['Net Floor FTE']:.2f}")
    m[1].metric("Client target", f"{rw['Client Target']:.0f}")
    m[2].metric("Buffered (min)", f"{rw['Buffered FTE'] * 30:+.0f}")
    m[3].metric("On lunch / short", f"{rw['On Lunch']:.1f} / {rw['On Short Break']:.1f}")
    ob = on_break(kk)
    if len(ob):
        st.dataframe(ob, hide_index=True)
    else:
        st.info("Nobody is on a break in this interval.")
    best = it[it["Scheduled FTE"] > 0].nlargest(3, "Buffered FTE")
    st.caption("Best intervals to move breaks into: " + ", ".join(f"{r.Interval} ({r['Buffered FTE'] * 30:+.0f} min)"
                                                                    for _, r in best.iterrows()))

# ----------------------------------------------------------------------------- Module 3
with tab3:
    st.caption(f"Basis: **{basis}** · Unit: **{unit}**. 🟩 surplus · 🟨 exact target · 🟥 understaffed · grey = closed")
    week, ot = {}, {}
    for j, d in enumerate(DATES):
        a = agg(vecs, DATES, j, mask_fn)
        f = basis_fte(a, basis)
        t_d = tgt_for(targets, sel_lobs, d)
        week[d] = np.where((t_d == 0) & (f == 0), np.nan, np.round(f - t_d, 2))
        ot[d] = float(np.clip(t_d - f, 0, None).sum() * 0.5)
    week = pd.DataFrame(week, index=SLOTS)
    ot = pd.Series(ot)
    f_sel = basis_fte(a_day, basis)
    deficit = np.clip(tgt - f_sel, 0, None)

    k = st.columns(4)
    k[0].metric("OT hours needed (week)", f"{ot.sum():.1f}")
    k[1].metric(f"OT hours needed ({sel_date})", f"{ot[sel_date]:.1f}")
    k[2].metric(f"Understaffed intervals ({sel_date})", int((deficit > 0).sum()))
    k[3].metric("Largest deficit (FTE)", f"{deficit.max():.1f}")

    st.subheader("Weekly 30-min heatmap (FTE vs target)")
    HFMT = "{:+.0f}" if unit == "Minutes" else "{:+.2f}"
    st.dataframe(smap((week * UF).style, heat_color).format(HFMT, na_rep="·"), height=620)

    st.subheader(f"{sel_date}: heatmap by LOB")
    cols = {}
    for lob in sel_lobs:
        a = agg(vecs, DATES, di, lambda d, lob=lob: mask_fn(d, [lob]))
        t, f = tgt_for(targets, [lob], sel_date), basis_fte(a, basis)
        cols[lob] = np.where((t == 0) & (f == 0), np.nan, np.round(f - t, 2))
    lob_heat = pd.DataFrame(cols, index=SLOTS)
    st.dataframe(smap((lob_heat * UF).style, heat_color).format(HFMT, na_rep="·"), height=400)

    c1, c2 = st.columns(2)
    c1.subheader("OT hours by day")
    c1.bar_chart(ot)
    c2.subheader(f"Understaffed intervals - {sel_date}")
    ot_df = pd.DataFrame({"Interval": SLOTS, "Target": tgt, "FTE": f_sel.round(2),
                          "Deficit": deficit.round(2), "OT Hours": (deficit * 0.5).round(2)})
    c2.dataframe(ot_df[ot_df["Deficit"] > 0], hide_index=True)

    with st.expander("🧑‍🤝‍🧑 OT & VTO suggestions", expanded=False):
        st.markdown("**Overtime: who to ask for each understaffed window**")
        ots = ot_suggestions(deficit)
        if len(ots):
            st.dataframe(ots, hide_index=True)
        else:
            st.success("No understaffed windows on this date.")
        st.markdown(f"**VTO: surplus windows (keeping {VTO_KEEP:g} FTE buffer, at least 2 hours)**")
        vto = vto_suggestions(it["Buffered FTE"].to_numpy())
        if len(vto):
            st.dataframe(vto, hide_index=True)
        else:
            st.info("No long surplus windows on this date.")

    d1, d2 = st.columns(2)
    d1.download_button("⬇️ Weekly interval matrix (CSV)", (week * UF).rename_axis("Interval").reset_index().to_csv(index=False)
                       .encode("utf-8-sig"), "weekly_interval_matrix.csv", "text/csv")
    d2.download_button("⬇️ OT analysis for date (CSV)", ot_df.to_csv(index=False).encode("utf-8-sig"),
                       f"ot_{sel_date}.csv", "text/csv")

# ----------------------------------------------------------------------------- Module 4
FMT4 = {"Shrinkage %": "{:.1f}%", "Attendance %": "{:.1f}%", "Night Allowance (EGP)": "{:,.0f}", "Total Scheduled": "{:.0f}",
        "Attended": "{:.0f}", "Unplanned Absences": "{:.0f}", "Late Shifts": "{:.0f}", "Overnight Shifts": "{:.0f}"}
SHOW4 = ["Total Scheduled", "Attended", "Unplanned Absences", "Shrinkage %", "Attendance %", "Late Shifts",
         "Overnight Shifts", "Night Allowance (EGP)"]
show, exc = pd.DataFrame(), pd.DataFrame()
with tab4:
    ad = agent_day(sel_date)
    if ad.empty:
        st.info("No scheduled agents for this selection.")
    else:
        g = summarize(ad, "LOB")
        T = g.iloc[-1]
        k = st.columns(4)
        k[0].metric("Scheduled FTE", int(T["Total Scheduled"]))
        k[1].metric("Attended FTE", int(T["Attended"]), delta=int(T["Attended"] - T["Total Scheduled"]))
        k[2].metric("Attendance %", f"{T['Attendance %']:.1f}%")
        k[3].metric("Shrinkage %", f"{T['Shrinkage %']:.1f}%", help="(Absence + late + early-leave hours) / scheduled hours")

        st.subheader(f"Executive summary by LOB - {sel_date}")
        show = g[["LOB"] + SHOW4]
        st.dataframe(show.style.format(FMT4), hide_index=True)
        st.caption("Allowance applies to attended agents only; rates and hours are configurable in the sidebar.")

        c1, c2 = st.columns(2)
        c1.subheader("By Team Leader")
        c1.dataframe(summarize(ad, "TL")[["TL"] + SHOW4].style.format(FMT4), hide_index=True)
        c2.subheader("Lost hours by exception type")
        by_exc = ad[ad["Exception"] != "None"].groupby("Exception")["Lost"].sum()
        if len(by_exc):
            c2.bar_chart(by_exc)
        else:
            c2.info("No exceptions logged on this date.")

        st.subheader("Weekly trend")
        frames = [x.assign(Date=d) for d in DATES if len(x := agent_day(d))]
        wk = pd.concat(frames)
        tr = wk.groupby("Date", sort=False).agg(Sched=("Employee ID", "count"), Att=("Attended", "sum"),
                                                Lost=("Lost", "sum"), Hrs=("Hours", "sum"))
        tr["Attendance %"] = tr["Att"] / tr["Sched"] * 100
        tr["Shrinkage %"] = tr["Lost"] / tr["Hrs"] * 100
        t1, t2 = st.columns(2)
        t1.line_chart(tr[["Attendance %", "Shrinkage %"]])
        wk_exc = wk[wk["Exception"] != "None"].groupby(["TL", "Exception"])["Lost"].sum().unstack(fill_value=0)
        if len(wk_exc):
            t2.caption("Lost hours this week by TL and exception type")
            t2.bar_chart(wk_exc)

        st.subheader("Exceptions logged")
        exc = ad[ad["Exception"] != "None"][["Employee ID", "Name", "LOB", "TL", "Shift", "Exception", "Exception Minutes"]]
        if len(exc):
            st.dataframe(exc, hide_index=True)
        else:
            st.success("No exceptions logged.")
        st.download_button("⬇️ Executive summary (CSV)", show.to_csv(index=False).encode("utf-8-sig"),
                           f"attendance_summary_{sel_date}.csv", "text/csv")

# ----------------------------------------------------------------------------- Module 5: roster editor
with tab5:
    st.subheader("Roster editor & shift swaps")
    opts = sorted(set(all_shifts) | set(SHIFTS) | {"OFF"},
                  key=lambda x: (SHIFTS.index(x) if x in SHIFTS else 98 if x == "OFF" else 99, x))
    rv = roster[roster["LOB"].isin(sel_lobs) & roster["TL"].isin(sel_tls)].reset_index(drop=True)
    edr = st.data_editor(rv, key=f"roster_ed_{hash((tuple(sel_lobs), tuple(sel_tls)))}", hide_index=True,
                         disabled=META, height=min(35 * len(rv) + 40, 520),
                         column_config={d: st.column_config.SelectboxColumn(d, options=opts, required=True) for d in DATES})
    if not edr[DATES].astype(str).equals(rv[DATES].astype(str)):
        new, log = roster.set_index("Employee ID").copy(), []
        for _, r in edr.iterrows():
            for d in DATES:
                old = new.at[r["Employee ID"], d]
                if str(r[d]) != str(old):
                    log.append({"Time": datetime.now().strftime("%Y-%m-%d %H:%M"), "Employee ID": r["Employee ID"],
                                "Name": new.at[r["Employee ID"], "Name"], "Date": d, "From": old, "To": r[d],
                                "Source": "Editor"})
                    new.at[r["Employee ID"], d] = r[d]
        apply_roster(new.reset_index()[META + DATES], log)

    st.markdown("**Swap shifts between two agents**")
    names = (roster["Employee ID"] + " - " + roster["Name"]).tolist()
    c1, c2, c3, c4 = st.columns([3, 3, 2, 1])
    ag_a = c1.selectbox("Agent A", names, key="swA")
    ag_b = c2.selectbox("Agent B", names, index=min(1, len(names) - 1), key="swB")
    sw_d = c3.selectbox("Date", DATES, index=di, key="swD")
    if c4.button("Swap"):
        ia, ib = names.index(ag_a), names.index(ag_b)
        if ia == ib:
            st.error("Pick two different agents.")
        else:
            new = roster.copy()
            va, vb = new.at[ia, sw_d], new.at[ib, sw_d]
            new.at[ia, sw_d], new.at[ib, sw_d] = vb, va
            now = datetime.now().strftime("%Y-%m-%d %H:%M")
            apply_roster(new, [
                {"Time": now, "Employee ID": roster.at[ia, "Employee ID"], "Name": roster.at[ia, "Name"], "Date": sw_d,
                 "From": va, "To": vb, "Source": f"Swap with {roster.at[ib, 'Name']}"},
                {"Time": now, "Employee ID": roster.at[ib, "Employee ID"], "Name": roster.at[ib, "Name"], "Date": sw_d,
                 "From": vb, "To": va, "Source": f"Swap with {roster.at[ia, 'Name']}"}])

    st.subheader("Roster rule checks")
    chk = roster_checks(roster)
    if len(chk):
        st.warning(f"{len(chk)} issue(s): minimum rest {MIN_REST}h between shifts, max {MAX_CONSEC} consecutive days.")
        st.dataframe(chk, hide_index=True)
    else:
        st.success("No rest-period or consecutive-day issues.")

    st.subheader("Change log")
    lg = pd.DataFrame(st.session_state.get("log", []))
    if len(lg):
        st.dataframe(lg, hide_index=True)
        st.download_button("⬇️ Change log (CSV)", lg.to_csv(index=False).encode("utf-8-sig"), "roster_change_log.csv", "text/csv")
    else:
        st.caption("No roster changes yet. Edits and swaps made here are logged and update the break tables automatically.")

# ----------------------------------------------------------------------------- save & export (sidebar, end of script)
with st.sidebar:
    st.divider()
    st.header("💾 Save & export")
    st.download_button("Save session (.xlsx)", session_bytes(), "wfm_session.xlsx", XLSX,
                       help="Roster, targets, all break edits and the change log. Restore it under Data inputs.")
    if st.button("Prepare Excel report"):
        st.session_state.report = make_report()
    if st.session_state.get("report"):
        st.download_button("⬇️ Download Excel report", st.session_state.report, f"wfm_report_{sel_date}.xlsx", XLSX)
